"""Build the summary .xlsx workbook from parsed Trip Generation Manual rows."""

from __future__ import annotations

from collections.abc import Iterable

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .parsing import Row
from .recommend import FITTED_CURVE, LINE_AT_CLUSTER, WEIGHTED_AVG, recommend_method

COLUMNS = [
    ("PDF Page", "pdf_page"),
    ("Volume", "volume"),
    ("Code", "code"),
    ("Land Use Name", "land_use_name"),
    ("Setting/Location", "setting"),
    ("Dependent Variable", "dependent_variable"),
    ("Independent Variable", "independent_variable"),
    ("Time Period", "time_period"),
    ("# of Studies", "n_studies"),
    ("Avg. Independent Variable", "avg_iv"),
    ("Directional Distribution", "dir_dist"),
    ("Average Rate", "avg_rate"),
    ("Range of Rates", "range_of_rates"),
    ("Standard Deviation", "stdev"),
    ("R²", "rsq"),
    ("Fitted Curve Equation", "equation"),
    ("Caution", "caution"),
    ("Recommended Method", "recommended_method"),
]

_COL_INDEX = {name: i + 1 for i, (name, _) in enumerate(COLUMNS)}
AVERAGE_RATE_COL = get_column_letter(_COL_INDEX["Average Rate"])
EQUATION_COL = get_column_letter(_COL_INDEX["Fitted Curve Equation"])
RECOMMENDED_COL = get_column_letter(_COL_INDEX["Recommended Method"])

_DEFAULT_WIDTHS = {
    "PDF Page": 9,
    "Volume": 9,
    "Code": 7,
    "Land Use Name": 45,
    "Setting/Location": 20,
    "Dependent Variable": 20,
    "Independent Variable": 24,
    "Time Period": 38,
    "# of Studies": 10,
    "Avg. Independent Variable": 14,
    "Directional Distribution": 16,
    "Average Rate": 12,
    "Range of Rates": 15,
    "Standard Deviation": 12,
    "R²": 8,
    "Fitted Curve Equation": 30,
    "Caution": 22,
    "Recommended Method": 18,
}

_HIGHLIGHT_FILL = PatternFill(fgColor="FF92D050", fill_type="solid")


def _write_header(ws: Worksheet) -> None:
    for col_idx, (name, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font = Font(bold=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = _DEFAULT_WIDTHS[name]
    ws.freeze_panes = "A2"


def _write_row(ws: Worksheet, row_idx: int, row: Row, method: str) -> None:
    values = {**row.__dict__, "recommended_method": method}
    for col_idx, (_, attr) in enumerate(COLUMNS, start=1):
        ws.cell(row=row_idx, column=col_idx, value=values[attr])


def _add_highlight_rules(ws: Worksheet, last_row: int) -> None:
    avg_range = f"{AVERAGE_RATE_COL}2:{AVERAGE_RATE_COL}{last_row}"
    eq_range = f"{EQUATION_COL}2:{EQUATION_COL}{last_row}"
    ws.conditional_formatting.add(
        avg_range,
        FormulaRule(
            formula=[f'${RECOMMENDED_COL}2="{WEIGHTED_AVG}"'],
            fill=_HIGHLIGHT_FILL,
        ),
    )
    for method in (FITTED_CURVE, LINE_AT_CLUSTER):
        ws.conditional_formatting.add(
            eq_range,
            FormulaRule(
                formula=[f'${RECOMMENDED_COL}2="{method}"'],
                fill=_HIGHLIGHT_FILL,
            ),
        )


def _write_index_sheet(ws: Worksheet, rows: list[Row]) -> None:
    ws.cell(row=1, column=1, value="Code").font = Font(bold=True)
    ws.cell(row=1, column=2, value="Land Use Name").font = Font(bold=True)
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 55

    seen: dict[tuple[int, str], None] = {}
    for row in rows:
        seen.setdefault((row.code, row.land_use_name), None)

    for i, (code, name) in enumerate(sorted(seen), start=2):
        ws.cell(row=i, column=1, value=code)
        ws.cell(row=i, column=2, value=name)
    ws.freeze_panes = "A2"


def build_workbook(rows: Iterable[Row]) -> Workbook:
    rows = list(rows)

    wb = Workbook()
    data_ws = wb.active
    data_ws.title = "Trip Generation Data"
    _write_header(data_ws)

    for i, row in enumerate(rows, start=2):
        method = recommend_method(
            n_studies=row.n_studies,
            avg_rate=row.avg_rate,
            stdev=row.stdev,
            rsq=row.rsq,
            equation=row.equation,
        )
        _write_row(data_ws, i, row, method)

    last_row = len(rows) + 1
    if len(rows) > 0:
        data_ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last_row}"
        _add_highlight_rules(data_ws, last_row)

    index_ws = wb.create_sheet("Land Use Index")
    _write_index_sheet(index_ws, rows)

    return wb
